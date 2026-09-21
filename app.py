import os
import sys
import io
import uuid
import time
import numpy as np
import cv2
from PIL import Image

import torch
import timm
import pymysql
from functools import wraps
from threading import Lock
from torchvision.models.detection import fasterrcnn_resnet50_fpn
from flask import (Flask, request, jsonify, render_template,
                   session, redirect, url_for, send_file)
from werkzeug.security import check_password_hash, generate_password_hash

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage

WEB_ROOT = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(WEB_ROOT)
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

import crops

CONF_THR = 0.5
MAX_CROPS = 6
NMS_IOU = 0.5
SIM_THRESHOLD = 0.6
CLASS_CN = {"scratches": "划痕(scratches)", "stain": "污渍(stain)"}

MODEL_PKG = os.path.join(PROJECT_ROOT, "weights", "afdn_model.pt")

app = Flask(__name__,
            template_folder=os.path.join(WEB_ROOT, "templates"),
            static_folder=os.path.join(WEB_ROOT, "static"))
app.secret_key = "afdn-web-secret-key-2026"

DB_CONFIG = {
    "host": "localhost",
    "port": 33061,
    "user": "root",
    "password": "root",
    "database": "AFDN",
    "charset": "utf8mb4",
}


def get_db():
    return pymysql.connect(**DB_CONFIG)


def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if not session.get("user"):
            return redirect(url_for("login"))
        return f(*args, **kwargs)
    return decorated


BATCH_RESULTS = {}
_BATCH_LOCK = Lock()
_BATCH_MAX = 50

device = None
rcnn = None
dinov2 = None
_dinov2_tf = None
proto = {}


def _load_dinov2(sd, device):
    from torchvision import transforms as T
    model = timm.create_model("vit_base_patch14_dinov2", pretrained=False,
                              num_classes=0, img_size=224)
    model.load_state_dict(sd, strict=True)
    model.eval().to(device)
    mean = (0.485, 0.456, 0.406)
    std = (0.229, 0.224, 0.225)

    def _letterbox(img, size=224):
        w, h = img.size
        scale = min(size / w, size / h)
        nw, nh = max(1, round(w * scale)), max(1, round(h * scale))
        img = img.resize((nw, nh), Image.BICUBIC)
        fill = tuple(int(round(m * 255)) for m in mean)
        canvas = Image.new("RGB", (size, size), fill)
        canvas.paste(img, ((size - nw) // 2, (size - nh) // 2))
        return canvas

    tf = T.Compose([T.Lambda(_letterbox), T.ToTensor(), T.Normalize(mean=mean, std=std)])
    return model, tf


@torch.no_grad()
def _extract(pil_img):
    import torch.nn.functional as F
    if not isinstance(pil_img, Image.Image):
        pil_img = Image.fromarray(pil_img)
    x = _dinov2_tf(pil_img.convert("RGB")).unsqueeze(0).to(device)
    emb = F.normalize(dinov2(x), dim=-1)
    return emb.squeeze(0).cpu().numpy().astype(np.float32)


print(">>> 正在加载模型包 afdn_model.pt（首次启动约需 10-20 秒）...")
device = torch.device("cuda" if torch.cuda.is_available() else "cpu")
print(f">>> Device: {device}")

if not os.path.isfile(MODEL_PKG):
    raise FileNotFoundError(
        f"模型包不存在: {MODEL_PKG}\n"
        "请先运行: py -3 build_model_package.py")

pkg = torch.load(MODEL_PKG, map_location="cpu")
print(f">>> 模型包版本: {pkg.get('version', 'unknown')}")

cfg = pkg.get("config", {})
CONF_THR = cfg.get("conf_thr", CONF_THR)
MAX_CROPS = cfg.get("max_crops", MAX_CROPS)
NMS_IOU = cfg.get("nms_iou", NMS_IOU)
SIM_THRESHOLD = cfg.get("sim_threshold", SIM_THRESHOLD)
print(f">>> 配置: conf={CONF_THR} max_crops={MAX_CROPS} nms_iou={NMS_IOU} sim_thr={SIM_THRESHOLD}")

rcnn = fasterrcnn_resnet50_fpn(weights=None, num_classes=pkg.get("rcnn_num_classes", 3))
rcnn.load_state_dict(pkg["rcnn_state_dict"])
rcnn.to(device).eval()
print(f">>> Faster R-CNN 加载完成")

dinov2, _dinov2_tf = _load_dinov2(pkg["dinov2_state_dict"], device)
print(f">>> DINOv2 加载完成")

proto = {k: v.numpy() for k, v in pkg["prototypes"].items()}
print(f">>> 原型就绪: scratches={proto['scratches'].shape}, stain={proto['stain'].shape}")
print(">>> 模型与原型就绪，可以开始上传图片判别。")


def _detect(bgr):
    rgb = cv2.cvtColor(bgr, cv2.COLOR_BGR2RGB)
    t = torch.from_numpy(rgb).permute(2, 0, 1).contiguous().to(device).float().div_(255.0)
    with torch.no_grad():
        preds = rcnn([t])[0]
    boxes = preds["boxes"].cpu().numpy()
    labels = preds["labels"].cpu().numpy().astype(np.int64)
    scores = preds["scores"].cpu().numpy().astype(np.float64)

    mask = scores > CONF_THR
    for i in range(len(labels)):
        if int(labels[i]) not in crops.CLASS_NAMES:
            mask[i] = False
    idx = np.where(mask)[0]
    if len(idx) == 0:
        return []
    boxes, labels, scores = boxes[idx], labels[idx], scores[idx]
    keep = crops._class_wise_nms(boxes, labels, scores, iou_thr=NMS_IOU)
    boxes, labels, scores = boxes[keep], labels[keep], scores[keep]
    order = np.argsort(-scores)[:MAX_CROPS]

    H, W = bgr.shape[:2]
    out = []
    for rank, k in enumerate(order, start=1):
        x1, y1, x2, y2 = boxes[k].tolist()
        nx1, ny1, nx2, ny2 = crops._expand_bbox(x1, y1, x2, y2, W, H, expand_ratio=0.08)
        w, h = nx2 - nx1, ny2 - ny1
        if w <= 16 or h <= 16:
            continue
        if max(w / max(1, h), h / max(1, w)) > 50:
            continue
        out.append({"rank": rank, "cls": crops.CLASS_NAMES[int(labels[k])],
                    "score": float(scores[k]), "box": [nx1, ny1, nx2, ny2]})
    return out


def _recommend_select(crop_embs):
    alive = list(range(len(crop_embs)))
    while len(alive) > 1:
        M = np.stack([crop_embs[i] for i in alive])
        mean = M.mean(axis=0)
        sims = [float(np.dot(crop_embs[i], mean) /
                      (np.linalg.norm(crop_embs[i]) * np.linalg.norm(mean) + 1e-8))
                for i in alive]
        alive.pop(int(np.argmin(sims)))
    return alive[0]


def infer(bgr):
    dets = _detect(bgr)
    crop_embs, crop_meta = [], []
    for d in dets:
        x1, y1, x2, y2 = d["box"]
        patch = bgr[max(0, y1):max(0, y2), max(0, x1):max(0, x2)]
        if patch.size == 0:
            continue
        pil = Image.fromarray(cv2.cvtColor(patch, cv2.COLOR_BGR2RGB))
        crop_embs.append(_extract(pil))
        crop_meta.append(d)

    if crop_embs:
        surv = _recommend_select(crop_embs)
        q = crop_embs[surv]
        sims = {k: float(np.dot(q, v)) for k, v in proto.items()}
        emb_cls = max(sims, key=sims.get)
        emb_sim = sims[emb_cls]
        rank1_cls = crop_meta[0]["cls"]
        det_available = True
        n_crops = len(crop_embs)
    else:
        pil = Image.fromarray(cv2.cvtColor(bgr, cv2.COLOR_BGR2RGB))
        q = _extract(pil)
        sims = {k: float(np.dot(q, v)) for k, v in proto.items()}
        emb_cls = max(sims, key=sims.get)
        emb_sim = sims[emb_cls]
        rank1_cls = None
        det_available = False
        n_crops = 0

    if not det_available:
        final, rule = emb_cls, "无检测框，纯特征匹配"
    elif emb_cls == rank1_cls:
        final, rule = emb_cls, "特征链与检测链一致"
    elif emb_sim >= SIM_THRESHOLD:
        final = emb_cls
        rule = f"两链冲突但特征相似度 {emb_sim:.2f} ≥ {SIM_THRESHOLD}，采信特征"
    else:
        final = rank1_cls
        rule = f"两链冲突且特征相似度 {emb_sim:.2f} < {SIM_THRESHOLD}，采信检测 rank-1"

    return {
        "category": final, "category_cn": CLASS_CN[final],
        "similarity": round(emb_sim, 4),
        "sim_scratches": round(sims["scratches"], 4),
        "sim_stain": round(sims["stain"], 4),
        "emb_category": emb_cls,
        "det_category": rank1_cls if det_available else "none",
        "n_crops": n_crops,
        "detections": [{"rank": d["rank"], "cls": d["cls"],
                        "score": round(d["score"], 4), "box": d["box"]}
                       for d in crop_meta],
        "rule": rule,
    }


def _draw(bgr, result):
    vis = bgr.copy()
    for d in result["detections"]:
        x1, y1, x2, y2 = d["box"]
        is_rank1 = d["rank"] == 1
        color = (0, 200, 0) if is_rank1 else (255, 160, 0)
        cv2.rectangle(vis, (x1, y1), (x2, y2), color, 4 if is_rank1 else 2)
        tag = f"#{d['rank']} {d['cls']} {d['score']:.2f}"
        cv2.putText(vis, tag, (x1, max(20, y1 - 8)),
                    cv2.FONT_HERSHEY_SIMPLEX, 0.8, color, 2)
    if not result["detections"]:
        cv2.putText(vis, "NO DETECTION (fallback: whole image)", (20, 40),
                    cv2.FONT_HERSHEY_SIMPLEX, 1.0, (0, 0, 255), 2)
    return vis


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        if not username or not password:
            return render_template("login.html", error="请输入用户名和密码")
        try:
            conn = get_db()
            cur = conn.cursor()
            cur.execute("SELECT password_hash FROM users WHERE username=%s", (username,))
            row = cur.fetchone()
            conn.close()
        except Exception as e:
            return render_template("login.html", error=f"数据库错误：{e}")
        if row and check_password_hash(row[0], password):
            session["user"] = username
            return redirect(url_for("index"))
        return render_template("login.html", error="用户名或密码错误")
    return render_template("login.html")


@app.route("/logout")
def logout():
    session.pop("user", None)
    return redirect(url_for("login"))


@app.route("/users")
@login_required
def users():
    conn = get_db()
    cur = conn.cursor()
    cur.execute("SELECT id, username, created_at FROM users ORDER BY id")
    user_list = cur.fetchall()
    conn.close()
    return render_template("users.html", users=user_list,
                           current_user=session.get("user"),
                           msg=request.args.get("msg"),
                           msg_type=request.args.get("msg_type"),
                           user=session.get("user"),
                           active="users")


@app.route("/users/add", methods=["POST"])
@login_required
def add_user():
    username = request.form.get("username", "").strip()
    password = request.form.get("password", "")
    if not username or not password:
        return redirect(url_for("users", msg="用户名和密码不能为空", msg_type="error"))
    try:
        conn = get_db()
        cur = conn.cursor()
        cur.execute("SELECT id FROM users WHERE username=%s", (username,))
        if cur.fetchone():
            conn.close()
            return redirect(url_for("users", msg="用户名已存在", msg_type="error"))
        cur.execute("INSERT INTO users (username, password_hash) VALUES (%s, %s)",
                    (username, generate_password_hash(password)))
        conn.commit()
        conn.close()
        return redirect(url_for("users", msg=f"用户 {username} 添加成功", msg_type="ok"))
    except Exception as e:
        return redirect(url_for("users", msg=f"添加失败：{e}", msg_type="error"))


@app.route("/users/<int:uid>/update", methods=["POST"])
@login_required
def update_user(uid):
    new_username = request.form.get("username", "").strip()
    new_password = request.form.get("password", "")
    if not new_username:
        return redirect(url_for("users", msg="用户名不能为空", msg_type="error"))
    try:
        conn = get_db()
        cur = conn.cursor()
        cur.execute("SELECT username FROM users WHERE id=%s", (uid,))
        old = cur.fetchone()
        if not old:
            conn.close()
            return redirect(url_for("users", msg="用户不存在", msg_type="error"))
        is_self = old[0] == session.get("user")
        cur.execute("SELECT id FROM users WHERE username=%s AND id<>%s",
                    (new_username, uid))
        if cur.fetchone():
            conn.close()
            return redirect(url_for("users", msg="用户名已存在", msg_type="error"))
        if new_password:
            cur.execute("UPDATE users SET username=%s, password_hash=%s WHERE id=%s",
                        (new_username, generate_password_hash(new_password), uid))
        else:
            cur.execute("UPDATE users SET username=%s WHERE id=%s",
                        (new_username, uid))
        conn.commit()
        if is_self:
            session["user"] = new_username
        conn.close()
        return redirect(url_for("users", msg="修改成功", msg_type="ok"))
    except Exception as e:
        return redirect(url_for("users", msg=f"修改失败：{e}", msg_type="error"))


@app.route("/users/<int:uid>/delete", methods=["POST"])
@login_required
def delete_user(uid):
    try:
        conn = get_db()
        cur = conn.cursor()
        cur.execute("SELECT username FROM users WHERE id=%s", (uid,))
        row = cur.fetchone()
        if not row:
            conn.close()
            return redirect(url_for("users", msg="用户不存在", msg_type="error"))
        if row[0] == session.get("user"):
            conn.close()
            return redirect(url_for("users", msg="不能删除当前登录的自己", msg_type="error"))
        cur.execute("DELETE FROM users WHERE id=%s", (uid,))
        conn.commit()
        conn.close()
        return redirect(url_for("users", msg=f"用户 {row[0]} 已删除", msg_type="ok"))
    except Exception as e:
        return redirect(url_for("users", msg=f"删除失败：{e}", msg_type="error"))


@app.route("/")
@login_required
def index():
    return render_template("index.html", user=session.get("user"), active="defect")


@app.route("/detect")
@login_required
def detect():
    return render_template("detect.html", user=session.get("user"), active="detect")


@app.route("/repair")
@login_required
def repair():
    return render_template("repair.html", user=session.get("user"), active="repair")


def _process_upload(f):
    try:
        data = np.frombuffer(f.read(), dtype=np.uint8)
        bgr = cv2.imdecode(data, cv2.IMREAD_COLOR)
        if bgr is None:
            return None, "图片解码失败，请上传 jpg/png 格式的有效图片"
    except Exception as e:
        return None, f"读取图片失败：{e}"
    try:
        result = infer(bgr)
    except Exception as e:
        return None, f"推理失败：{e}"

    result["filename"] = os.path.basename(f.filename or "未命名.jpg")
    vis = _draw(bgr, result)
    out_name = f"{uuid.uuid4().hex}.jpg"
    out_dir = os.path.join(WEB_ROOT, "static", "results")
    os.makedirs(out_dir, exist_ok=True)
    vis_path = None
    ok, buf = cv2.imencode(".jpg", vis)
    if ok:
        vis_path = os.path.join(out_dir, out_name)
        buf.tofile(vis_path)
        result["result_image"] = f"/static/results/{out_name}"
    return {"result": result, "vis_path": vis_path}, None


@app.route("/predict", methods=["POST"])
@login_required
def predict():
    f = request.files.get("image")
    if f is None or f.filename == "":
        return jsonify({"error": "未收到图片文件"}), 400
    item, err = _process_upload(f)
    if err:
        return jsonify({"error": err}), 400
    return jsonify(item["result"])


@app.route("/batch/start", methods=["POST"])
@login_required
def batch_start():
    bid = uuid.uuid4().hex
    with _BATCH_LOCK:
        BATCH_RESULTS[bid] = {"created": time.time(), "results": []}
        if len(BATCH_RESULTS) > _BATCH_MAX:
            oldest = min(BATCH_RESULTS.items(), key=lambda kv: kv[1]["created"])[0]
            BATCH_RESULTS.pop(oldest, None)
    return jsonify({"batch_id": bid})


@app.route("/batch/<bid>/add", methods=["POST"])
@login_required
def batch_add(bid):
    with _BATCH_LOCK:
        batch = BATCH_RESULTS.get(bid)
    if batch is None:
        return jsonify({"error": "批次不存在或已过期，请刷新页面后重试"}), 404
    f = request.files.get("image")
    if f is None or f.filename == "":
        return jsonify({"error": "未收到图片文件"}), 400
    item, err = _process_upload(f)
    if err:
        return jsonify({"error": err, "filename": os.path.basename(f.filename or "")}), 400
    with _BATCH_LOCK:
        batch["results"].append(item)
    return jsonify(item["result"])


@app.route("/export_excel/<bid>")
@login_required
def export_excel(bid):
    with _BATCH_LOCK:
        batch = BATCH_RESULTS.get(bid)
        items = list(batch["results"]) if batch else None
    if not items:
        return "批次不存在、已过期或没有可导出的结果", 404

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "判别结果"

    headers = ["序号", "文件名", "最终类别", "检出缺陷框数", "融合规则", "检测可视化"]
    widths = [6, 22, 16, 12, 46, 38]

    head_fill = PatternFill("solid", fgColor="2B6CB0")
    head_font = Font(color="FFFFFF", bold=True, size=11)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="CBD5E0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    scratch_fill = PatternFill("solid", fgColor="EBF4FF")
    stain_fill = PatternFill("solid", fgColor="FFF5F5")

    for c, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(1, c, h)
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = center
        cell.border = border
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w
    ws.freeze_panes = "A2"

    for i, item in enumerate(items, start=2):
        r = item["result"]
        row_fill = scratch_fill if r["category"] == "scratches" else stain_fill
        vals = [
            i - 1, r.get("filename", ""), r["category_cn"],
            r["n_crops"], r["rule"],
        ]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(i, c, v)
            cell.border = border
            cell.fill = row_fill
            cell.alignment = left if c in (2, 5) else center

        vis_path = item.get("vis_path")
        if vis_path and os.path.isfile(vis_path):
            try:
                pim = Image.open(vis_path)
                w0, h0 = pim.size
                tw = 260
                th = max(1, int(h0 * tw / w0))
                pim = pim.resize((tw, th), Image.BICUBIC)
                buf = io.BytesIO()
                pim.save(buf, format="JPEG", quality=85)
                buf.seek(0)
                xim = XLImage(buf)
                ws.add_image(xim, f"F{i}")
                ws.row_dimensions[i].height = th * 0.76
            except Exception:
                pass

    ws.row_dimensions[1].height = 24
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    fname = f"AFDN判别结果_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(
        out, as_attachment=True, download_name=fname,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000, debug=False, threaded=True)
